"""Resume text extraction (PDF / DOCX / TXT / MD)."""
import os

from .textutil import normalize_ws


class ResumeError(Exception):
    pass


def extract_text(path):
    if not path:
        raise ResumeError("no resume path given")
    path = os.path.expanduser(path)
    if not os.path.exists(path):
        raise ResumeError(f"resume not found: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext in (".docx", ".docm"):
        return _docx(path)
    if ext in (".txt", ".md", ".markdown", ".rtf", ".html", ".htm"):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            raw = f.read()
        if ext in (".html", ".htm"):
            from .textutil import html_to_text
            return html_to_text(raw)
        return normalize_ws(raw)
    if ext == ".doc":
        raise ResumeError("legacy .doc is not supported; save the resume as .docx or .pdf")
    # unknown extension: try as text
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return normalize_ws(f.read())


def _pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ResumeError("pypdf is missing: run  python -m pip install --user pypdf") from e
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    text = normalize_ws(" ".join(parts))
    if len(text) < 80:
        raise ResumeError("could not extract text from the PDF (scanned image?). Export a text-based PDF or DOCX.")
    return text


def _docx(path):
    try:
        import docx
    except ImportError as e:
        raise ResumeError("python-docx is missing: run  python -m pip install --user python-docx") from e
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return normalize_ws(" ".join(parts))
