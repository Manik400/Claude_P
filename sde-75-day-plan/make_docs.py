"""Word and Excel versions of plan.json, for printing and offline tracking.

    python make_docs.py        -> dist/75-Day-SDE-Plan-Final.docx, dist/75-Day-SDE-Plan-Final.xlsx

The Excel file is a working tracker: tick ✓ in the Tracker sheet and the
Dashboard, the day status (Completed / Partial / Missed / Today / Upcoming)
and the LeetCode sheet update by formula.
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
DIST = HERE / "dist"
GREEN, GREEN_SOFT = "15803D", "E7F3EC"
AMBER, RED, BLUE, PURPLE, GREY = "B45309", "B91C1C", "1D4ED8", "7C3AED", "6B6560"
DIFF_COLOR = {"Easy": GREEN, "Medium": AMBER, "Hard": RED}
TICK = "✓"


def load():
    return json.loads((HERE / "plan.json").read_text(encoding="utf-8"))


def nice(d: str) -> str:
    x = date.fromisoformat(d)
    return f"{x.strftime('%a')} {x.day} {x.strftime('%b %Y')}"


# ================================================================ Word

def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def hyperlink(paragraph, url, text, bold=False, color=BLUE):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    if bold:
        rpr.append(OxmlElement("w:b"))
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def run(p, text, bold=False, color=None, size=None, italic=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size)
    return r


def table(doc, header, rows, widths=None, head_fill=GREEN):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run(c.paragraphs[0], h, bold=True, color="FFFFFF", size=9.5)
        shade(c, head_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run(cells[i].paragraphs[0], str(v), size=9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def prob_line(cell, x):
    """One SQL / design question: tick box, linked name, difficulty, and the spec when it is a build-it task."""
    para = cell.paragraphs[0]
    run(para, "☐  ", size=10)
    if x.get("url"):
        hyperlink(para, x["url"], x["name"], bold=True)
    else:
        run(para, x["name"], bold=True, size=10)
    if not x.get("review"):
        run(para, f"   {x['difficulty']}", bold=True, color=DIFF_COLOR[x["difficulty"]], size=9)
        if x.get("source"):
            run(para, f"   {x['source']}", size=9, color=GREY)
    if x.get("statement"):
        run(cell.add_paragraph(), x["statement"], size=9, color=GREY)


def make_docx(plan) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for name, size, color in (("Heading 1", 18, GREEN), ("Heading 2", 14, "1C1917"), ("Heading 3", 12, "1C1917")):
        s = doc.styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(color)

    # ---- cover
    p = doc.add_paragraph()
    run(p, plan["title"], bold=True, size=28, color=GREEN)
    p = doc.add_paragraph()
    run(p, plan["subtitle"], size=12, color=GREY)
    p = doc.add_paragraph()
    run(p, f"{nice(plan['start'])}  →  {nice(plan['end'])}   ·   75 days   ·   225 LeetCode   ·   3 phases", bold=True, size=11)
    p = doc.add_paragraph()
    run(p, "Day 1 is a Monday so every Build Saturday and Review Sunday falls on a real weekend. "
           "Tick tasks in the Excel tracker or in the Plan tab of the phone site.", size=10, color=GREY, italic=True)

    doc.add_heading("How this plan works", 1)
    for r in plan["rules"]:
        doc.add_paragraph(r, style="List Bullet")
    doc.add_heading("Daily non-negotiables", 2)
    for h in plan["habits"]:
        doc.add_paragraph("☐  " + h["label"])
    doc.add_paragraph("☐  3 LeetCode    ☐  HLD read    ☐  LLD in C#    ☐  Cloud lab")
    doc.add_heading("Easy day (the floor on a drained day)", 2)
    doc.add_paragraph("  ·  ".join(plan["easy_day"]))

    doc.add_heading("Timetable", 1)
    for key, title in (("weekday", "Monday to Friday"), ("saturday", "Saturday - build day"), ("sunday", "Sunday - review day")):
        doc.add_heading(title, 3)
        table(doc, ["Time", "What", "Notes"], plan["schedule"][key], widths=[2.6, 7, 7.4])

    doc.add_heading("The three phases", 1)
    table(doc, ["Phase", "Days", "Dates", "What it covers"],
          [[f"{ph['id']} · {ph['name']}", f"{ph['days'][0]}–{ph['days'][1]}",
            f"{nice(plan['days'][ph['days'][0] - 1]['date'])} – {nice(plan['days'][ph['days'][1] - 1]['date'])}", ph["summary"]]
           for ph in plan["phases"]], widths=[3.6, 1.6, 4.8, 7])

    doc.add_heading("DSA patterns - the one-line approach", 2)
    table(doc, ["Pattern", "Approach"], [[k, v] for k, v in plan["patterns"].items()], widths=[4.6, 12.4])

    # ---- day by day
    for ph in plan["phases"]:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(f"Phase {ph['id']} · {ph['name']} (days {ph['days'][0]}–{ph['days'][1]})", 1)
        doc.add_paragraph(ph["summary"]).runs[0].italic = True
        week = None
        for d in plan["days"][ph["days"][0] - 1: ph["days"][1]]:
            if d["week"] != week:
                week = d["week"]
                doc.add_heading(f"Week {week}", 2)
            h = doc.add_heading(level=3)
            run(h, f"Day {d['day']} · {nice(d['date'])} · ", color=GREY)
            run(h, d["theme"])
            p = doc.add_paragraph()
            run(p, d["type"], bold=True, color=PURPLE if "Sunday" in d["type"] else AMBER if "Saturday" in d["type"] else GREEN, size=9.5)
            run(p, f"   ·   {d['cloud']['provider']}   ·   Pattern: {d['dsa']['pattern']}", size=9.5, color=GREY)

            t = doc.add_table(rows=0, cols=2)
            t.style = "Table Grid"

            def row(label, fill):
                cells = t.add_row().cells
                cells[0].width, cells[1].width = Cm(3.2), Cm(13.8)
                cells[0].text = ""
                run(cells[0].paragraphs[0], label, bold=True, size=9.5)
                shade(cells[0], fill)
                cells[1].text = ""
                return cells[1]

            c = row("LeetCode\n7:50 am 2 Medium\n10:15 pm 1 Hard", "E7F3EC")
            for i, q in enumerate(d["leetcode"]):
                para = c.paragraphs[0] if i == 0 else c.add_paragraph()
                run(para, "☐  ", size=10)
                hyperlink(para, q["url"], q["name"], bold=True)
                run(para, f"   {q['difficulty']}", bold=True, color=DIFF_COLOR[q["difficulty"]], size=9)
            para = c.add_paragraph()
            run(para, "Approach: ", bold=True, size=9)
            run(para, d["dsa"]["tip"], size=9, color=GREY)

            c = row("SQL\nlunch, 15 min", "E7F3EC")
            prob_line(c, d["sql"])

            c = row("HLD read\nlunch, 25 min", "E8EEFB")
            run(c.paragraphs[0], "☐  " + d["hld"]["topic"], bold=True, size=10)
            for pt in d["hld"]["points"]:
                run(c.add_paragraph(), "•  " + pt, size=9.5, color=GREY)

            c = row("LLD · C#\n9:00 pm, 30 min", "F1ECFD")
            run(c.paragraphs[0], "☐  " + d["lld"]["topic"], bold=True, size=10)
            run(c.add_paragraph(), "Done when: " + d["lld"]["deliverable"], size=9, color=GREY)

            c = row("Design coding\n9:30 pm, 20 min", "E8EEFB")
            prob_line(c, d["design"])

            c = row(f"Cloud · {d['cloud']['provider']}\n9:50 pm, 25 min", "FBF1E6")
            run(c.paragraphs[0], "☐  " + d["cloud"]["topic"], bold=True, size=10)
            run(c.add_paragraph(), d["cloud"]["note"], size=9, color=GREY)

            c = row("Habits", "FBEAEA")
            run(c.paragraphs[0], "   ".join("☐ " + x for x in ("Walk 1 km", "Water 2.5–3 L", "Healthy food", "Chill 1 h", "Sleep 11:30")), size=9.5)

            if d.get("checkpoint"):
                c = row("CHECKPOINT\npass criteria", "FFF4D6")
                for i, crit in enumerate(d["checkpoint"]):
                    para = c.paragraphs[0] if i == 0 else c.add_paragraph()
                    run(para, "☐  " + crit, size=9.5, bold=True)
            doc.add_paragraph()

    # ---- LeetCode index
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("All 225 LeetCode problems", 1)
    lt = table(doc, ["#", "Day", "Date", "Problem", "Difficulty", "Pattern", "Solved"], [], widths=[1, 1.1, 2.6, 6.4, 1.8, 3.4, 1.2])
    n = 0
    for d in plan["days"]:
        for q in d["leetcode"]:
            n += 1
            cells = lt.add_row().cells
            for i, v in enumerate([str(n), str(d["day"]), nice(d["date"])[:-5]]):
                cells[i].text = ""
                run(cells[i].paragraphs[0], v, size=9)
            cells[3].text = ""
            hyperlink(cells[3].paragraphs[0], q["url"], q["name"])
            cells[4].text = ""
            run(cells[4].paragraphs[0], q["difficulty"], bold=True, color=DIFF_COLOR[q["difficulty"]], size=9)
            cells[5].text = ""
            run(cells[5].paragraphs[0], q["pattern"], size=9)
            cells[6].text = ""
            run(cells[6].paragraphs[0], "☐", size=10)

    # ---- SQL and design question lists
    for key, title in (("sql", "All SQL questions"), ("design", "All design coding questions")):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(title, 1)
        st = table(doc, ["Day", "Date", "Question", "Difficulty", "Done"], [], widths=[1.1, 2.6, 10.2, 2, 1.1])
        for d in plan["days"]:
            x = d[key]
            cells = st.add_row().cells
            for i, v in enumerate([str(d["day"]), nice(d["date"])[:-5]]):
                cells[i].text = ""
                run(cells[i].paragraphs[0], v, size=9)
            cells[2].text = ""
            if x.get("url"):
                hyperlink(cells[2].paragraphs[0], x["url"], x["name"])
            else:
                run(cells[2].paragraphs[0], x["name"], size=9.5, italic=bool(x.get("review")))
            if x.get("statement"):
                run(cells[2].add_paragraph(), x["statement"], size=8.5, color=GREY)
            cells[3].text = ""
            run(cells[3].paragraphs[0], x["difficulty"], bold=True, color=DIFF_COLOR.get(x["difficulty"], GREY), size=9)
            cells[4].text = ""
            run(cells[4].paragraphs[0], "☐", size=10)

    # ---- tracker grid
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("75-day tracker", 1)
    doc.add_paragraph("LC = 3 LeetCode · SQL · Dsg = design question · HLD · LLD · Cloud · W = walk · H2O = water · F = food · C = chill · S = sleep")
    table(doc, ["Day", "Date", "LC", "SQL", "Dsg", "HLD", "LLD", "Cloud", "W", "H2O", "F", "C", "S", "Notes"],
          [[d["day"], nice(d["date"])[:-5]] + ["☐"] * 11 + [""] for d in plan["days"]],
          widths=[0.9, 2.2, 0.8, 0.8, 0.8, 0.8, 0.8, 1, 0.7, 0.8, 0.7, 0.7, 0.7, 4.5])

    DIST.mkdir(exist_ok=True)
    out = DIST / "75-Day-SDE-Plan-Final.docx"
    doc.save(out)
    return out


def make_qna_docx(plan, qna) -> Path:
    """Every day's HLD and LLD Q&A, grouped by level, for reading away from the screen."""
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    p = doc.add_paragraph()
    run(p, "75-Day SDE Plan - HLD & LLD Q&A", bold=True, size=24, color=GREEN)
    p = doc.add_paragraph()
    run(p, f"{sum(len(v['hld']) + len(v['lld']) for v in qna.values())} interview questions, 6 HLD + 6 LLD per day, "
           "grouped Basic / Intermediate / Advanced. Say the Basic answers out loud; try the Advanced ones before reading.", size=11, color=GREY)
    LEVEL_COLOR = {"Basic": GREEN, "Intermediate": AMBER, "Advanced": RED}
    for d in plan["days"]:
        v = qna.get(str(d["day"]))
        if not v:
            continue
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        h = doc.add_heading(level=1)
        run(h, f"Day {d['day']} · {nice(d['date'])} · ", color=GREY, size=15)
        run(h, d["theme"], size=15)
        for kind, label, topic in (("hld", "HLD", d["hld"]["topic"]), ("lld", "LLD", d["lld"]["topic"])):
            h = doc.add_heading(level=2)
            run(h, f"{label}: {topic}", size=12.5)
            level = None
            for i, item in enumerate(v[kind], 1):
                if item["level"] != level:
                    level = item["level"]
                    para = doc.add_paragraph()
                    run(para, level.upper(), bold=True, size=9, color=LEVEL_COLOR.get(level, GREY))
                para = doc.add_paragraph()
                run(para, f"Q{i}. ", bold=True, color=GREEN)
                run(para, item["q"], bold=True)
                answer_paragraphs(doc, item["a"])
    DIST.mkdir(exist_ok=True)
    out = DIST / "75-Day-SDE-QnA.docx"
    doc.save(out)
    return out


def answer_paragraphs(doc, text):
    """The answer text with its "\n- " bullets and ```csharp``` blocks laid out."""
    import re as _re
    parts = _re.split(r"```(?:csharp|cs|c#)?\n?", text)
    for i, chunk in enumerate(parts):
        if i % 2:
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.6)
            r = para.add_run(chunk.rstrip("\n"))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue
        chunk = chunk.strip()
        if not chunk:
            continue
        if chunk.startswith("- "):
            chunk = "\n" + chunk
        lines = chunk.split("\n- ")
        if lines[0].strip():
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.6)
            run(para, lines[0].strip(), size=10)
        for line in lines[1:]:
            para = doc.add_paragraph(style="List Bullet")
            para.paragraph_format.left_indent = Cm(1.2)
            run(para, line.strip(), size=10)


# ================================================================ Excel

HEAD_FILL = PatternFill("solid", fgColor=GREEN)
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="E6E2DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


def header(ws, names, widths, freeze="C2"):
    ws.append(names)
    for i, (n, w) in enumerate(zip(names, widths), 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center", vertical="center", wrap_text=True), BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    # set once: openpyxl leaves a stale pane selection behind when this is
    # changed afterwards, and Excel then refuses to open the file
    ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 30


def make_xlsx(plan) -> Path:
    wb = Workbook()
    days = plan["days"]
    N = len(days)

    # ---------------- Tracker (the one place you tick)
    tr = wb.active
    tr.title = "Tracker"
    # E-G LC, H SQL, I Design, J HLD, K LLD, L Cloud, M Walk, N Water, O Food, P Chill, Q Sleep, R Done, S Of, T %, U Status, V Notes
    cols = ["Day", "Date", "Type", "Theme", "LC 1 (M)", "LC 2 (M)", "LC 3 (H)", "SQL", "Design", "HLD", "LLD", "Cloud",
            "Walk", "Water (L)", "Food", "Chill", "Sleep", "Done", "Of", "%", "Status", "Notes"]
    header(tr, cols, [6, 13, 15, 30, 7, 7, 7, 6, 7, 6, 6, 7, 6, 8, 6, 6, 6, 6, 5, 7, 12, 30])
    for i, d in enumerate(days, start=2):
        tr.append([d["day"], date.fromisoformat(d["date"]), d["type"], d["theme"]] + [""] * 13)
        tr.cell(row=i, column=2).number_format = "ddd d mmm"
        # done = ticks + water goal met
        tr.cell(row=i, column=18, value=f'=COUNTIF(E{i}:M{i},"{TICK}")+COUNTIF(O{i}:Q{i},"{TICK}")+IF(N(N{i})>=2.5,1,0)')
        tr.cell(row=i, column=19, value=13)
        tr.cell(row=i, column=20, value=f"=R{i}/S{i}")
        tr.cell(row=i, column=20).number_format = "0%"
        # Completed > Today > Partial (some ticks) > Upcoming (future, none) > Missed (past, none)
        tr.cell(row=i, column=21, value=(f'=IF(R{i}>=S{i},"Completed",IF(B{i}=TODAY(),"Today",IF(R{i}>0,"Partial",'
                                         f'IF(B{i}>TODAY(),"Upcoming","Missed"))))'))
        for c in range(1, 23):
            cell = tr.cell(row=i, column=c)
            cell.border = BORDER
            if 5 <= c <= 21:
                cell.alignment = CENTER
    tick_dv = DataValidation(type="list", formula1=f'"{TICK}"', allow_blank=True)
    tick_dv.prompt, tick_dv.promptTitle = f"Pick {TICK} when done", "Done?"
    tr.add_data_validation(tick_dv)
    tick_dv.add(f"E2:M{N + 1}")
    tick_dv.add(f"O2:Q{N + 1}")
    water_dv = DataValidation(type="list", formula1='"0,0.5,1,1.5,2,2.5,3"', allow_blank=True)
    tr.add_data_validation(water_dv)
    water_dv.add(f"N2:N{N + 1}")
    rng = f"U2:U{N + 1}"
    for text, fill, font in (("Completed", "DCF3E4", GREEN), ("Partial", "FDEBD3", AMBER), ("Missed", "FBE0E0", RED),
                             ("Today", "15803D", "FFFFFF"), ("Upcoming", "F0EEE9", GREY)):
        tr.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{text}"'],
                                                      fill=PatternFill("solid", fgColor=fill), font=Font(bold=True, color=font)))
    tr.conditional_formatting.add(f"E2:Q{N + 1}", CellIsRule(operator="equal", formula=[f'"{TICK}"'],
                                                             fill=PatternFill("solid", fgColor="DCF3E4"), font=Font(bold=True, color=GREEN)))
    tr.conditional_formatting.add(f"A2:D{N + 1}", FormulaRule(formula=[f"$B2=TODAY()"], fill=PatternFill("solid", fgColor="E7F3EC"),
                                                              font=Font(bold=True)))
    tr.auto_filter.ref = f"A1:V{N + 1}"

    # ---------------- Plan (read-only detail)
    pl = wb.create_sheet("Plan")
    header(pl, ["Day", "Date", "Weekday", "Type", "Phase", "Week", "Theme", "DSA pattern", "Approach",
                "LeetCode 1", "Diff", "LeetCode 2", "Diff", "LeetCode 3", "Diff",
                "HLD read (lunch)", "HLD key points", "LLD in C# (9 pm)", "Cloud lab (9:50 pm)", "Checkpoint pass criteria",
                "SQL (lunch)", "Diff", "Design coding (9:30 pm)", "Diff"],
           [6, 12, 8, 14, 7, 6, 26, 20, 40, 30, 8, 30, 8, 30, 8, 36, 60, 40, 44, 50, 36, 8, 40, 8])
    for i, d in enumerate(days, start=2):
        lc = d["leetcode"]
        pl.append([d["day"], date.fromisoformat(d["date"]), d["weekday"], d["type"], d["phase"], d["week"], d["theme"],
                   d["dsa"]["pattern"], d["dsa"]["tip"],
                   lc[0]["name"], lc[0]["difficulty"], lc[1]["name"], lc[1]["difficulty"], lc[2]["name"], lc[2]["difficulty"],
                   d["hld"]["topic"], "\n".join("• " + p for p in d["hld"]["points"]), d["lld"]["topic"],
                   f"{d['cloud']['provider']}: {d['cloud']['topic']}", "\n".join("• " + c for c in d["checkpoint"] or []),
                   d["sql"]["name"], d["sql"]["difficulty"],
                   d["design"]["name"] + (("\n" + d["design"]["statement"]) if d["design"].get("statement") else ""), d["design"]["difficulty"]])
        for col, x in ((21, d["sql"]), (23, d["design"])):
            if x.get("url"):
                pl.cell(row=i, column=col).hyperlink = x["url"]
                pl.cell(row=i, column=col).font = Font(color=BLUE, underline="single")
            pl.cell(row=i, column=col + 1).font = Font(bold=True, color=DIFF_COLOR.get(x["difficulty"], GREY))
        pl.cell(row=i, column=2).number_format = "ddd d mmm"
        for k, col in enumerate((10, 12, 14)):
            c = pl.cell(row=i, column=col)
            c.hyperlink = lc[k]["url"]
            c.font = Font(color=BLUE, underline="single")
            dc = pl.cell(row=i, column=col + 1)
            dc.font = Font(bold=True, color=DIFF_COLOR[lc[k]["difficulty"]])
        for c in range(1, 25):
            pl.cell(row=i, column=c).alignment = WRAP
            pl.cell(row=i, column=c).border = BORDER
        pl.row_dimensions[i].height = 90 if d["design"].get("statement") else 62 if d["checkpoint"] else 48
    pl.auto_filter.ref = f"A1:X{N + 1}"

    # ---------------- LeetCode (solved comes from the Tracker ticks)
    lcs = wb.create_sheet("LeetCode")
    header(lcs, ["#", "Day", "Date", "Problem", "Difficulty", "Pattern", "Solved", "Revisit", "Notes"],
           [5, 6, 12, 46, 10, 30, 9, 9, 40], freeze="E2")
    n = 1
    for d in days:
        for k, q in enumerate(d["leetcode"]):
            n += 1
            tick_col = get_column_letter(5 + k)          # Tracker E/F/G
            lcs.append([n - 1, d["day"], date.fromisoformat(d["date"]), q["name"], q["difficulty"], q["pattern"],
                        f'=IF(Tracker!{tick_col}{d["day"] + 1}="{TICK}","Yes","")', ""])
            lcs.cell(row=n, column=3).number_format = "d mmm"
            c = lcs.cell(row=n, column=4)
            c.hyperlink, c.font = q["url"], Font(color=BLUE, underline="single")
            lcs.cell(row=n, column=5).font = Font(bold=True, color=DIFF_COLOR[q["difficulty"]])
            for col in range(1, 10):
                lcs.cell(row=n, column=col).border = BORDER
    rv = DataValidation(type="list", formula1='"↻"', allow_blank=True)
    lcs.add_data_validation(rv)
    rv.add(f"H2:H{n}")
    lcs.conditional_formatting.add(f"G2:G{n}", CellIsRule(operator="equal", formula=['"Yes"'],
                                                          fill=PatternFill("solid", fgColor="DCF3E4"), font=Font(bold=True, color=GREEN)))
    lcs.auto_filter.ref = f"A1:I{n}"

    # ---------------- Dashboard
    db = wb.create_sheet("Dashboard", 0)
    db.column_dimensions["A"].width = 34
    db.column_dimensions["B"].width = 16
    db.column_dimensions["C"].width = 16
    db.column_dimensions["D"].width = 16
    db["A1"] = plan["title"]
    db["A1"].font = Font(bold=True, size=20, color=GREEN)
    db["A2"] = f"{nice(plan['start'])} → {nice(plan['end'])}  ·  tick {TICK} in the Tracker sheet; everything here updates"
    db["A2"].font = Font(color=GREY, italic=True)
    last = N + 1
    rows = [
        ("Overall progress (tasks)", f"=SUM(Tracker!R2:R{last})/SUM(Tracker!S2:S{last})", "0%"),
        ("Days completed", f'=COUNTIF(Tracker!U2:U{last},"Completed")', "0"),
        ("Days partial", f'=COUNTIF(Tracker!U2:U{last},"Partial")', "0"),
        ("Days missed", f'=COUNTIF(Tracker!U2:U{last},"Missed")', "0"),
        ("Today is plan day", f'=IFERROR(MATCH(TODAY(),Tracker!B2:B{last},0),"-")', "0"),
        ("LeetCode solved (of 225)", f'=COUNTIF(LeetCode!G2:G{n},"Yes")', "0"),
        ("LeetCode marked ↻ revisit", f'=COUNTIF(LeetCode!H2:H{n},"↻")', "0"),
        ("SQL questions done (of 75)", f'=COUNTIF(Tracker!H2:H{last},"{TICK}")', "0"),
        ("Design questions done (of 75)", f'=COUNTIF(Tracker!I2:I{last},"{TICK}")', "0"),
        ("HLD reads done", f'=COUNTIF(Tracker!J2:J{last},"{TICK}")', "0"),
        ("LLD builds done", f'=COUNTIF(Tracker!K2:K{last},"{TICK}")', "0"),
        ("Cloud labs done", f'=COUNTIF(Tracker!L2:L{last},"{TICK}")', "0"),
        ("Walk / exercise days", f'=COUNTIF(Tracker!M2:M{last},"{TICK}")', "0"),
        ("Water goal days (2.5 L+)", f'=COUNTIF(Tracker!N2:N{last},">=2.5")', "0"),
        ("Q&A marked known (of 900)", f'=COUNTIF(\'Q&A\'!G2:G901,"{TICK}")', "0"),
    ]
    r = 4
    for label, f, fmt in rows:
        db.cell(row=r, column=1, value=label).font = Font(bold=True)
        c = db.cell(row=r, column=2, value=f)
        c.number_format, c.font, c.alignment = fmt, Font(bold=True, size=13, color=GREEN), CENTER
        r += 1
    r += 1
    db.cell(row=r, column=1, value="Phase").font = HEAD_FONT
    for i, h in enumerate(("Phase", "Days", "Progress", "Completed days"), 1):
        c = db.cell(row=r, column=i, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for ph in plan["phases"]:
        r += 1
        a, b = ph["days"][0] + 1, ph["days"][1] + 1
        db.cell(row=r, column=1, value=f"{ph['id']} · {ph['name']}")
        db.cell(row=r, column=2, value=f"{ph['days'][0]}–{ph['days'][1]}")
        c = db.cell(row=r, column=3, value=f"=SUM(Tracker!R{a}:R{b})/SUM(Tracker!S{a}:S{b})")
        c.number_format = "0%"
        db.cell(row=r, column=4, value=f'=COUNTIF(Tracker!U{a}:U{b},"Completed")')
    db.conditional_formatting.add(f"C{r - 2}:C{r}", _databar())

    # ---------------- Q&A (tick G when you can answer without looking)
    qa = wb.create_sheet("Q&A")
    header(qa, ["Day", "Date", "Section", "Level", "Question", "Answer", "Known"], [6, 12, 9, 13, 60, 90, 8], freeze="E2")
    r = 1
    lv_color = {"Basic": GREEN, "Intermediate": AMBER, "Advanced": RED}
    for d in days:
        v = QNA.get(str(d["day"])) or {}
        for kind in ("hld", "lld"):
            for item in v.get(kind, []):
                r += 1
                qa.append([d["day"], date.fromisoformat(d["date"]), kind.upper(), item["level"], item["q"], item["a"], ""])
                qa.cell(row=r, column=2).number_format = "d mmm"
                qa.cell(row=r, column=4).font = Font(bold=True, color=lv_color.get(item["level"], GREY))
                for col in range(1, 8):
                    qa.cell(row=r, column=col).alignment = WRAP
                    qa.cell(row=r, column=col).border = BORDER
    kdv = DataValidation(type="list", formula1=f'"{TICK}"', allow_blank=True)
    qa.add_data_validation(kdv)
    kdv.add(f"G2:G{max(r, 2)}")
    qa.auto_filter.ref = f"A1:G{max(r, 2)}"

    # ---------------- Timetable + rules
    tt = wb.create_sheet("Timetable")
    header(tt, ["When", "Time", "What", "Notes"], [14, 12, 48, 52], freeze="A2")
    for key, title in (("weekday", "Mon–Fri"), ("saturday", "Saturday"), ("sunday", "Sunday")):
        for t, what, note in plan["schedule"][key]:
            tt.append([title, t, what, note])
    tt.append([])
    tt.append(["Rules"])
    tt.cell(row=tt.max_row, column=1).font = Font(bold=True, color=GREEN, size=12)
    for rule in plan["rules"]:
        tt.append(["", "", rule])
        tt.cell(row=tt.max_row, column=3).alignment = WRAP
    tt.append(["Easy day", "", "  ·  ".join(plan["easy_day"])])

    DIST.mkdir(exist_ok=True)
    out = DIST / "75-Day-SDE-Plan-Final.xlsx"
    wb.save(out)
    return out


def _databar():
    from openpyxl.formatting.rule import DataBarRule
    return DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=GREEN)


QNA = {}


def main():
    global QNA
    plan = load()
    qpath = HERE / "qna.json"
    QNA = json.loads(qpath.read_text(encoding="utf-8")) if qpath.exists() else {}
    print(make_docx(plan))
    print(make_xlsx(plan))
    if QNA:
        print(make_qna_docx(plan, QNA))


if __name__ == "__main__":
    main()
