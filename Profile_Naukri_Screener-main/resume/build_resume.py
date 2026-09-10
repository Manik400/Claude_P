"""Generate an ATS-safe resume as .docx plus a plain-text version.

The content comes from `resume/resume.yaml`, which is gitignored - a resume is
personal data and does not belong in a shared repository. Copy
`resume.example.yaml` to `resume.yaml` and fill it in; this file only decides
how it is laid out.

ATS parsers read a Word file as a linear stream of paragraphs. Anything that
breaks that stream - tables, text boxes, multi-column layouts, headers and
footers, images, icon fonts - either scrambles the reading order or is dropped
outright. An earlier version of this resume lost its project labels exactly
that way: they parsed *after* their own descriptions.

So everything here is single-column body paragraphs. Section rules are drawn as
paragraph borders rather than table edges, dates sit inline rather than in a
right-hand column, and bullets are literal characters with a hanging indent
instead of Word list numbering, which some parsers strip.
"""
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

OUT = Path(__file__).resolve().parent
CONTENT = OUT / "resume.yaml"
EXAMPLE = OUT / "resume.example.yaml"


class ResumeError(RuntimeError):
    """resume.yaml is missing or does not have what the layout needs."""


def load(path: Path = CONTENT) -> dict:
    """Read resume.yaml and check it has enough to build a resume from.

    The checks are deliberately loud. A missing `experience` key would
    otherwise produce a clean-looking one-page resume with no jobs on it, and
    that is the kind of thing you notice after sending it somewhere.
    """
    if not path.exists():
        raise ResumeError(
            f"No {path.name} in {path.parent}.\n"
            f"  Copy the template and fill it in:\n"
            f"      cp resume/{EXAMPLE.name} resume/{path.name}")
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except yaml.YAMLError as exc:
        raise ResumeError(f"{path.name} is not valid YAML: {exc}")
    if not isinstance(data, dict):
        raise ResumeError(f"{path.name} must be a YAML mapping.")

    missing = [k for k in ("name", "title", "summary", "experience") if not data.get(k)]
    if missing:
        raise ResumeError(
            f"{path.name} is missing: {', '.join(missing)}. "
            f"See {EXAMPLE.name} for what each field holds.")

    for index, job in enumerate(data["experience"], start=1):
        absent = [k for k in ("title", "org", "dates", "bullets") if not job.get(k)]
        if absent:
            raise ResumeError(
                f"{path.name}: experience entry {index} "
                f"({job.get('title') or job.get('org') or 'unnamed'}) "
                f"is missing: {', '.join(absent)}")
    return data


def _pairs(rows: list, first: str, second: str) -> list[tuple[str, str]]:
    """Accept either a mapping or a two-item list per row.

    The example file uses mappings because they are self-documenting, but a
    plain `- [Languages, "Python, SQL"]` is less to type and works too.
    """
    out = []
    for row in rows or []:
        if isinstance(row, dict):
            out.append((str(row.get(first) or ""), str(row.get(second) or "")))
        elif isinstance(row, (list, tuple)) and len(row) >= 2:
            out.append((str(row[0]), str(row[1])))
    return out


INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def style_base(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"          # metric-safe and present on every parser's box
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.05

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)


def bottom_rule(paragraph) -> None:
    """Underline a heading with a paragraph border, not a table edge."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    borders.append(bottom)
    p_pr.append(borders)


def heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = ACCENT
    bottom_rule(p)


def bullet(doc: Document, text: str) -> None:
    """Literal bullet with a hanging indent - Word list numbering is often stripped."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.18)
    pf.first_line_indent = Inches(-0.18)
    pf.space_after = Pt(2)
    p.add_run("•  " + text)


def basename(data: dict) -> str:
    """File stem for the outputs. Derived from the name unless one is given."""
    stated = str(data.get("output_basename") or "").strip()
    if stated:
        return stated
    slug = "_".join(part.capitalize() for part in str(data["name"]).split())
    return f"{slug}_Resume" if slug else "Resume"


def build(data: dict) -> Path:
    doc = Document()
    style_base(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(data["name"])
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = ACCENT

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(data["title"])
    run.font.size = Pt(10.5)

    for line in data.get("contact") or []:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(str(line)).font.size = Pt(9.5)

    heading(doc, "Professional Summary")
    doc.add_paragraph(data["summary"])

    skills = _pairs(data.get("skills"), "label", "items")
    if skills:
        heading(doc, "Technical Skills")
        for label, items in skills:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"{label}: ").bold = True
            p.add_run(items)

    heading(doc, "Professional Experience")
    for job in data["experience"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        run = p.add_run(job["title"])
        run.bold = True
        run.font.size = Pt(10.5)

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        # `loc` is optional - a fully remote role has no office to name, and an
        # empty segment would leave a stray " |  | " in the line.
        run = p.add_run(" | ".join(str(bit) for bit in
                                   (job["org"], job.get("loc"), job["dates"]) if bit))
        run.italic = True
        run.font.size = Pt(9.5)

        for line in job["bullets"]:
            bullet(doc, line)

    projects = _pairs(data.get("projects"), "label", "description")
    if projects:
        heading(doc, "Key Projects")
        for label, desc in projects:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(0.18)
            pf.first_line_indent = Inches(-0.18)
            pf.space_after = Pt(3)
            p.add_run("•  ")
            p.add_run(f"{label}: ").bold = True   # label FIRST - an earlier version lost this
            p.add_run(desc)

    if data.get("certifications"):
        heading(doc, "Certifications & Continuous Learning")
        for line in data["certifications"]:
            bullet(doc, str(line))

    education = data.get("education") or []
    if education:
        heading(doc, "Education")
        for entry in education:
            if isinstance(entry, dict):
                degree, school, years = (entry.get("degree"), entry.get("institution"),
                                         entry.get("years"))
            else:
                degree, school, years = (list(entry) + [None, None])[:3]
            p = doc.add_paragraph()
            p.add_run(str(degree or "")).bold = True
            p.add_run("".join(f" | {bit}" for bit in (school, years) if bit))

    path = OUT / f"{basename(data)}.docx"
    doc.save(path)
    return path


def build_text(data: dict) -> Path:
    """Plain text for portals that mangle uploads and want a paste instead.

    This is also the file the interview-prep module reads, and its ALL-CAPS
    headings are load-bearing there: naukri/interview/profile.py splits on them
    to tell resume narrative (strong evidence for a skill) from a skills list
    (weaker). Renaming a heading here quietly downgrades your own evidence.
    """
    lines = [data["name"], data["title"], *(str(c) for c in data.get("contact") or []),
             "", "PROFESSIONAL SUMMARY", data["summary"]]

    skills = _pairs(data.get("skills"), "label", "items")
    if skills:
        lines += ["", "TECHNICAL SKILLS"]
        lines += [f"{label}: {items}" for label, items in skills]

    lines += ["", "PROFESSIONAL EXPERIENCE"]
    for job in data["experience"]:
        lines += ["", job["title"],
                  " | ".join(str(bit) for bit in
                             (job["org"], job.get("loc"), job["dates"]) if bit)]
        lines += [f"- {b}" for b in job["bullets"]]

    projects = _pairs(data.get("projects"), "label", "description")
    if projects:
        lines += ["", "KEY PROJECTS"]
        lines += [f"- {label}: {desc}" for label, desc in projects]

    if data.get("certifications"):
        lines += ["", "CERTIFICATIONS & CONTINUOUS LEARNING"]
        lines += [f"- {c}" for c in data["certifications"]]

    education = data.get("education") or []
    if education:
        lines += ["", "EDUCATION"]
        for entry in education:
            if isinstance(entry, dict):
                bits = (entry.get("degree"), entry.get("institution"), entry.get("years"))
            else:
                bits = tuple(entry)
            lines.append(" | ".join(str(b) for b in bits if b))

    path = OUT / f"{basename(data)}.txt"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def main() -> int:
    try:
        data = load()
    except ResumeError as exc:
        print(f"\n  {exc}\n")
        return 2
    docx_path = build(data)
    txt_path = build_text(data)
    words = len(txt_path.read_text(encoding="utf-8").split())
    print(f"wrote {docx_path}")
    print(f"wrote {txt_path}")
    print(f"word count: {words}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
