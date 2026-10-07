"""Resume text extraction (PDF / DOCX / TXT / MD)."""
import os

from .textutil import normalize_ws


class ResumeError(Exception):
    pass


# ----------------------------------------------------------------------------- the device's resume
#
# One resume per computer, outside the repo and outside any login:
# %LOCALAPPDATA%\JobHuntPhone\resume\current.<ext> (+ current.json: name, size, when). Every run
# on this PC - the worldwide search, the hourly rounds, company-site applies, the Premium runner -
# uses it unless a --resume is given on the command line. Another PC, or another Windows user on
# this one, has its own folder and so its own resume. Set it with site\set_resume.bat (drag a
# file onto it) or  python site\tools\resume_store.py set <file>.

DEVICE_RESUME_EXTS = (".pdf", ".docx", ".txt", ".md")


def device_store_dir():
    base = os.environ.get("JOBHUNT_DEVICE_DIR") or os.environ.get("LOCALAPPDATA") or os.path.expanduser("~")
    return os.path.join(base, "JobHuntPhone", "resume")


def device_resume_path():
    """The resume stored for this computer, or None."""
    d = device_store_dir()
    for ext in DEVICE_RESUME_EXTS:
        p = os.path.join(d, "current" + ext)
        if os.path.exists(p):
            return p
    return None


def device_resume_info():
    """{"path", "name", "size", "set_at"} for the stored resume, or None."""
    import json
    p = device_resume_path()
    if not p:
        return None
    info = {"path": p, "name": os.path.basename(p), "size": os.path.getsize(p), "set_at": ""}
    try:
        with open(os.path.join(device_store_dir(), "current.json"), encoding="utf-8") as f:
            info.update({k: v for k, v in json.load(f).items() if k in ("name", "set_at")})
    except (OSError, ValueError):
        pass
    return info


def set_device_resume(path):
    """Copy `path` into the device store as the current resume (the old one is replaced). Returns the stored path."""
    import json
    import shutil
    from datetime import datetime
    path = os.path.expanduser(path)
    ext = os.path.splitext(path)[1].lower()
    if ext not in DEVICE_RESUME_EXTS:
        raise ResumeError("a .pdf, .docx, .txt or .md resume, not " + (ext or "a file without an extension"))
    if not os.path.exists(path):
        raise ResumeError("resume not found: " + path)
    extract_text(path)          # unreadable files are refused before anything is replaced
    d = device_store_dir()
    os.makedirs(d, exist_ok=True)
    clear_device_resume()
    dest = os.path.join(d, "current" + ext)
    shutil.copyfile(path, dest)
    with open(os.path.join(d, "current.json"), "w", encoding="utf-8") as f:
        json.dump({"name": os.path.basename(path), "size": os.path.getsize(dest),
                   "set_at": datetime.now().isoformat(timespec="seconds"), "from": os.path.abspath(path)}, f, indent=1)
    return dest


def clear_device_resume():
    d = device_store_dir()
    for ext in DEVICE_RESUME_EXTS + (".json",):
        try:
            os.remove(os.path.join(d, "current" + ext))
        except OSError:
            pass


def extract_text(path):
    if not path:
        raise ResumeError("no resume path given")
    path = os.path.expanduser(path)
    if not os.path.exists(path):
        raise ResumeError(f"resume not found: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext in (".docx", ".docm"):
        return _docx(path)
    if ext in (".txt", ".md", ".markdown", ".rtf", ".html", ".htm"):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            raw = f.read()
        if ext in (".html", ".htm"):
            from .textutil import html_to_text
            return html_to_text(raw)
        return normalize_ws(raw)
    if ext == ".doc":
        raise ResumeError("legacy .doc is not supported; save the resume as .docx or .pdf")
    # unknown extension: try as text
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return normalize_ws(f.read())


def _pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ResumeError("pypdf is missing: run  python -m pip install --user pypdf") from e
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    text = normalize_ws(" ".join(parts))
    if len(text) < 80:
        raise ResumeError("could not extract text from the PDF (scanned image?). Export a text-based PDF or DOCX.")
    return text


def _docx(path):
    try:
        import docx
    except ImportError as e:
        raise ResumeError("python-docx is missing: run  python -m pip install --user python-docx") from e
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return normalize_ws(" ".join(parts))
